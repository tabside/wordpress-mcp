"""Configuration parsing, safe downloads and the Word conversion."""

from __future__ import annotations

import io

import httpx
import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.shared import OxmlElement, qn
from docx.shared import RGBColor

from wordpressmcp.config import ConfigError, Settings
from wordpressmcp.crypto import hash_password, verify_password
from wordpressmcp.docx_convert import DocxError, docx_to_content
from wordpressmcp.fetch import FetchError, fetch_url


# ---------------------------------------------------------------- config


def test_password_hash_roundtrip():
    encoded = hash_password("ein langes passwort", iterations=1000)
    assert "$" not in encoded and verify_password("ein langes passwort", encoded)
    assert not verify_password("falsch", encoded)
    assert not verify_password("x", "kaputt")


def test_sites_from_env(base_env, tmp_path):
    secret = tmp_path / "pw"
    secret.write_text("file secret\n")
    base_env.setenv("WP_SITE_ZWEITE_SEITE_URL", "https://www.example.org/blog/")
    base_env.setenv("WP_SITE_ZWEITE_SEITE_APP_PASSWORD_FILE", str(secret))
    base_env.setenv("WP_SITE_ZWEITE_SEITE_USER", "mcp-bot")
    base_env.setenv("WP_SITE_KUNDE_REST_URL", "https://kunde.example.com/?rest_route=")
    s = Settings.load()
    assert s.sites["kunde"].rest_url == "https://kunde.example.com/?rest_route="
    assert set(s.sites) == {"kunde", "zweite-seite"}
    site = s.sites["zweite-seite"]
    assert site.rest_url == "https://www.example.org/blog/wp-json"
    assert site.app_password == "file secret" and site.username == "mcp-bot" and site.label == "www.example.org"
    assert "file secret" not in repr(site)
    assert s.sites["kunde"].username == "claude-mcp"


def test_site_requires_https_and_password(base_env):
    base_env.setenv("WP_SITE_KUNDE_URL", "http://kunde.example.com")
    with pytest.raises(ConfigError, match="https"):
        Settings.load()
    base_env.setenv("WP_SITE_KUNDE_URL", "https://kunde.example.com")
    base_env.delenv("WP_SITE_KUNDE_APP_PASSWORD")
    with pytest.raises(ConfigError, match="app_password"):
        Settings.load()


def test_sites_from_json(base_env, tmp_path):
    base_env.setenv("WP_SITES", '[{"id": "Zwei", "url": "https://www.example.net", "app_password": "aaaa bbbb", "label": "Zwei GmbH"}]')
    s = Settings.load()
    assert set(s.sites) == {"kunde", "zwei"} and s.sites["zwei"].label == "Zwei GmbH" and s.sites["zwei"].app_password == "aaaa bbbb"
    base_env.setenv("WP_SITES", '[{"id": "kunde", "url": "https://www.example.net", "app_password": "x"}]')
    with pytest.raises(ConfigError, match="twice"):
        Settings.load()
    base_env.setenv("WP_SITES", '[{"id": "x", "url": "https://www.example.net", "password": "x"}]')
    with pytest.raises(ConfigError, match="unknown field"):
        Settings.load()
    base_env.setenv("WP_SITES", "{kaputt")
    with pytest.raises(ConfigError, match="JSON"):
        Settings.load()
    path = tmp_path / "sites.json"
    path.write_text('[{"id": "datei", "url": "https://www.example.net", "app_password": "y"}]')
    base_env.delenv("WP_SITES")
    base_env.setenv("WP_SITES_FILE", str(path))
    assert "datei" in Settings.load().sites


def test_no_sites_is_allowed(base_env):
    base_env.delenv("WP_SITE_KUNDE_URL")
    base_env.delenv("WP_SITE_KUNDE_APP_PASSWORD")
    assert Settings.load().sites == {}


def test_users_required_and_validated(base_env):
    base_env.setenv("GATEWAY_USERS", "kim=klartext")
    with pytest.raises(ConfigError, match="kim"):
        Settings.load()
    base_env.delenv("GATEWAY_USERS")
    with pytest.raises(ConfigError, match="GATEWAY_USERS"):
        Settings.load()


# ----------------------------------------------------------------- fetch


def _resolver(mapping):
    def resolve(host, port):
        return mapping[host]

    return resolve


@pytest.mark.anyio
async def test_fetch_blocks_private_addresses():
    for url, addresses in (
        ("http://intern.example.com/x.docx", ["10.0.0.5"]),
        ("http://intern.example.com/x.docx", ["93.184.216.34", "127.0.0.1"]),
        ("http://169.254.169.254/latest", None),
        ("http://[::1]/x", None),
    ):
        resolver = _resolver({"intern.example.com": addresses}) if addresses else _resolver({})
        with pytest.raises(FetchError, match="nicht öffentliche"):
            await fetch_url(url, max_bytes=1000, resolver=resolver, transport=httpx.MockTransport(lambda r: httpx.Response(200)))
    with pytest.raises(FetchError, match="http"):
        await fetch_url("file:///etc/passwd", max_bytes=1000)
    with pytest.raises(FetchError, match="Zugangsdaten"):
        await fetch_url("https://u:p@example.com/x", max_bytes=1000)


@pytest.mark.anyio
async def test_fetch_pins_ip_checks_redirects_and_size():
    seen = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append((request.url.host, request.headers["host"], request.extensions.get("sni_hostname")))
        if request.url.path == "/start":
            return httpx.Response(302, headers={"location": "https://files.example.com/doc.docx"})
        if request.url.path == "/doc.docx":
            return httpx.Response(200, content=b"x" * 10, headers={"content-type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"})
        if request.url.path == "/evil":
            return httpx.Response(302, headers={"location": "http://intern.example.com/admin"})
        return httpx.Response(200, content=b"y" * 5000)

    resolver = _resolver({"www.example.com": ["93.184.216.34"], "files.example.com": ["93.184.216.35"], "intern.example.com": ["192.168.1.1"]})
    out = await fetch_url("https://www.example.com/start", max_bytes=100, resolver=resolver, transport=httpx.MockTransport(handler))
    assert out.data == b"x" * 10 and out.filename == "doc.docx"
    assert seen == [("93.184.216.34", "www.example.com", "www.example.com"), ("93.184.216.35", "files.example.com", "files.example.com")]
    with pytest.raises(FetchError, match="nicht öffentliche"):
        await fetch_url("https://www.example.com/evil", max_bytes=100, resolver=resolver, transport=httpx.MockTransport(handler))
    with pytest.raises(FetchError, match="grösser"):
        await fetch_url("https://www.example.com/big", max_bytes=100, resolver=resolver, transport=httpx.MockTransport(handler))


@pytest.fixture
def anyio_backend():
    return "asyncio"


# ------------------------------------------------------------------ docx


def _link(par, url, text):
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    par._p.append(h)


def sample_docx() -> bytes:
    d = Document()
    d.add_heading("Reorganisation der Gemeinden", level=1)
    d.add_paragraph("")
    p = d.add_paragraph("Einleitung mit ")
    p.add_run("fett").bold = True
    p.add_run(" und ")
    p.add_run("kursiv").italic = True
    red = p.add_run(" rot")
    red.font.color.rgb = RGBColor(255, 0, 0)
    p.add_run(" und Link ")
    _link(p, "https://www.example.com/gesetz", "Gesetz")
    _link(p, "javascript:alert(1)", "böse")
    d.add_heading("Ausgangslage", level=2)
    d.add_paragraph("Punkt A", style="List Bullet")
    d.add_paragraph("Punkt B", style="List Bullet")
    d.add_paragraph("Schritt 1", style="List Number")
    d.add_heading("Details", level=3)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Frist"
    t.cell(0, 1).text = "30 Tage"
    d.add_paragraph("Ende gut.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_to_blocks():
    out = docx_to_content(sample_docx())
    c = out["content"]
    assert out["title"] == "Reorganisation der Gemeinden"
    assert "Reorganisation" not in c
    assert '<!-- wp:paragraph -->\n<p>Einleitung mit <strong>fett</strong> und <em>kursiv</em> rot und Link <a href="https://www.example.com/gesetz">Gesetz</a>böse</p>' in c
    assert '<!-- wp:heading -->\n<h2 class="wp-block-heading">Ausgangslage</h2>\n<!-- /wp:heading -->' in c
    assert '<!-- wp:heading {"level":3} -->\n<h3 class="wp-block-heading">Details</h3>' in c
    assert '<ul class="wp-block-list"><!-- wp:list-item -->\n<li>Punkt A</li>' in c
    assert '<!-- wp:list {"ordered":true} -->\n<ol class="wp-block-list">' in c
    assert '<figure class="wp-block-table"><table class="has-fixed-layout"><tbody><tr><td>Frist</td><td>30 Tage</td></tr></tbody></table></figure>' in c
    assert "Ende gut." in c and "\xa0" not in c
    assert "javascript" not in c and "style=" not in c and "<span" not in c
    assert "<p></p>" not in c
    assert out["stats"]["links"] == 1 and any("Link" in w for w in out["warnings"])


def test_docx_to_html_and_heading_shift():
    out = docx_to_content(sample_docx(), output="html", extract_title=False)
    assert out["content"].startswith("<h2>Reorganisation der Gemeinden</h2>")
    assert "<h3>Ausgangslage</h3>" in out["content"] and "<!-- wp:" not in out["content"]
    assert out["heading_shift"] == 1


def nested_list_docx() -> bytes:
    """Bullet list with a second level, numbered like Word does (numPr with ilvl)."""
    import copy

    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    d = Document()
    num_id = d.styles["List Bullet"].element.pPr.numPr.numId.val
    numbering = d.part.numbering_part.element
    num = [n for n in numbering.findall(w + "num") if n.get(w + "numId") == str(num_id)][0]
    abstract_id = num.find(w + "abstractNumId").get(w + "val")
    abstract = [a for a in numbering.findall(w + "abstractNum") if a.get(w + "abstractNumId") == abstract_id][0]
    abstract.find(w + "multiLevelType").set(w + "val", "hybridMultilevel")
    level1 = copy.deepcopy(abstract.find(w + "lvl"))
    level1.set(w + "ilvl", "1")
    level1.remove(level1.find(w + "pStyle"))
    abstract.append(level1)
    for text, level in (("Oben", 0), ("Unten A", 1), ("Unten B", 1), ("Oben 2", 0)):
        p = d.add_paragraph(text)
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = level
        num_pr.get_or_add_numId().val = num_id
    d.add_paragraph("Ein Zitat", style="Quote")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_nested_list_and_quote():
    c = docx_to_content(nested_list_docx(), extract_title=False)["content"]
    assert '<li>Oben<!-- wp:list -->\n<ul class="wp-block-list"><!-- wp:list-item -->\n<li>Unten A</li>' in c
    assert "<!-- /wp:list --></li>\n<!-- /wp:list-item -->" in c
    assert '<!-- wp:quote -->\n<blockquote class="wp-block-quote"><!-- wp:paragraph -->\n<p>Ein Zitat</p>' in c


def test_docx_rejects_other_files():
    with pytest.raises(DocxError, match="kein .docx"):
        docx_to_content(b"%PDF-1.4")
    buf = io.BytesIO()
    import zipfile

    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(DocxError, match="kein Word"):
        docx_to_content(buf.getvalue())


def test_docx_tool(gateway, access):
    import base64

    from conftest import call

    client, _, _ = gateway
    res = call(client, access, "docx_to_content", {"data_base64": base64.b64encode(sample_docx()).decode()})
    out = res["structuredContent"]
    assert out["title"] == "Reorganisation der Gemeinden" and out["format"] == "blocks"
    res = call(client, access, "docx_to_content", {})
    assert res["isError"] is True
